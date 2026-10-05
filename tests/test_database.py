from io import BytesIO

from openpyxl import load_workbook

from backend.database import analyze, facets, search
from backend.models import AnalysisRequest, Filters, SearchRequest


def test_prefix_accents_and_punctuation(sample_data):
    for term in ['metalur', 'Metalúrgica', 'METALÚR', 'metalúr***']:
        result = search(SearchRequest(filtros=Filters(termo=term)))
        assert result['total'] == 1
        assert result['itens'][0]['nome'] == 'Metalúrgica Sertão'
    assert search(SearchRequest(filtros=Filters(termo='***')))['total'] == 0


def test_combined_filters_case_and_opening_year(sample_data):
    filters = Filters(cidade='FORTALEZA', segmento='PADARIA E CONFEITARIA', bairro='aldeota', somente_celular=True, somente_email=True, somente_sem_dominio=True, score_minimo=80, ano_maximo=2020)
    result = search(SearchRequest(filtros=filters))
    assert result['total'] == 1
    assert result['itens'][0]['nome'] == 'Padaria São José'
    assert search(SearchRequest(filtros=Filters(cidade='juazeiro do norte')))['total'] == 1


def test_pagination_is_stable_and_complete(sample_data):
    pages = [search(SearchRequest(pagina=p, por_pagina=2)) for p in [1, 2, 3, 4]]
    assert [p['total'] for p in pages] == [5] * 4
    ids = [row['id'] for p in pages for row in p['itens']]
    assert len(ids) == len(set(ids)) == 5
    assert not pages[-1]['itens']


def test_facets_ignore_only_their_own_filter(sample_data):
    result = facets(Filters(cidade='Fortaleza', segmento='Padaria e confeitaria', somente_celular=True))
    assert {r['nome']: r['contatos'] for r in result['municipios']} == {'Fortaleza': 2, 'Juazeiro Do Norte': 1}
    assert {r['nome']: r['contatos'] for r in result['segmentos']} == {'Padaria e confeitaria': 2}


def test_aggregation_percentages_and_comparisons(sample_data):
    total = analyze(AnalysisRequest())[0]
    assert total['contatos'] == 5
    assert total['sem_dominio'] == 3  # A empresa sem e-mail não é contada.
    assert total['percentual_com_celular'] == 60
    rows = analyze(AnalysisRequest(filtros=Filters(cidades=['Fortaleza', 'Sobral']), agrupar_por=['cidade']))
    assert {r['cidade']: r['contatos'] for r in rows} == {'Fortaleza': 3, 'Sobral': 1}


def excel_cnpjs(content):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    assert workbook.sheetnames == ['Contatos']
    rows = list(workbook['Contatos'].values)
    columns = rows[2]
    assert columns[:5] == ('Nome', 'E-mail', 'Telefone', 'Telegram', 'Observações')
    values = [row[columns.index('CNPJ')] for row in rows[3:]]
    workbook.close()
    return values


def test_export_matches_every_filter_and_selection_takes_precedence(client):
    filters = {'cidade': 'Fortaleza', 'somente_celular': True, 'ano_minimo': 2024, 'score_minimo': 40}
    found = client.post('/api/contatos/buscar', json={'filtros': filters}).json()
    export = client.post('/api/exportar', json={'filtros': filters})
    assert export.status_code == 200
    assert excel_cnpjs(export.content) == [c['cnpj'] for c in found['itens']]
    selected = client.post('/api/exportar', json={'filtros': filters, 'ids': [1, 4]})
    assert set(excel_cnpjs(selected.content)) == {'11111111000100', '44444444000100'}


def test_invalid_inputs_other_states_and_origin_are_rejected(client):
    for body in [{'uf': 'SP'}, {'score_minimo': -1}, {'ordem': 'DROP TABLE contatos'}, {'ano_minimo': 2025, 'ano_maximo': 2020}]:
        assert client.post('/api/contatos/buscar', json={'filtros': body}).status_code == 422
    assert client.post('/api/analises', json={'agrupar_por': ['cnpj']}).status_code == 422
    assert client.post('/api/contatos/buscar', json={}, headers={'Origin': 'https://example.com'}).status_code == 403
    assert client.get('/api/base', headers={'Host': 'untrusted.example'}).status_code == 403
    assert client.get('/api/base', headers={'X-Forwarded-Host': 'untrusted.example'}).status_code == 403
    assert client.get('/api/malhas/malha_35.geojson').status_code == 404
    assert client.get('/api/contatos/9999').status_code == 404


def test_excel_treats_cadastral_values_as_text(client, sample_data):
    import sqlite3
    with sqlite3.connect(sample_data / 'uf' / 'CE' / 'contatos.db') as conn:
        conn.execute("UPDATE contatos SET nome='=1+1' WHERE id=1")
    exported = client.post('/api/exportar', json={'ids': [1]})
    workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=False)
    assert workbook['Contatos']['A4'].value == '=1+1'
    assert workbook['Contatos']['A4'].data_type == 's'
    workbook.close()


def test_export_rejects_a_stale_selection(client):
    result = client.post('/api/exportar', json={'ids': [1], 'base_gerada_em': '2025-01-01 00:00'})
    assert result.status_code == 409
    assert 'refaça a seleção' in result.json()['detail']
